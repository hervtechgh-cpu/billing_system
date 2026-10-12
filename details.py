import re
import pandas as pd
from pypdf import PdfReader
from openpyxl import load_workbook
from openpyxl.styles import Font
import tkinter as tk
from tkinter import filedialog, messagebox


# ============================================================
# SELECT PDF FROM LOCAL COMPUTER
# ============================================================

root = tk.Tk()
root.withdraw()

pdf_file = filedialog.askopenfilename(
    title="Select PDF File",
    filetypes=[
        ("PDF files", "*.pdf"),
        ("All files", "*.*")
    ]
)

if not pdf_file:
    print("No PDF file selected.")
    exit()

print(f"\nSelected PDF:")
print(pdf_file)


# ============================================================
# SELECT WHERE TO SAVE EXCEL FILE
# ============================================================

excel_file = filedialog.asksaveasfilename(
    title="Save Excel File As",
    defaultextension=".xlsx",
    filetypes=[
        ("Excel files", "*.xlsx")
    ]
)

if not excel_file:
    print("No output file selected.")
    exit()

print(f"\nExcel output:")
print(excel_file)


# ============================================================
# READ PDF
# ============================================================

print("\nReading PDF...")

reader = PdfReader(pdf_file)

all_text = ""

for page_number, page in enumerate(reader.pages, start=1):

    text = page.extract_text()

    if text:
        all_text += "\n" + text

    print(
        f"Processing page "
        f"{page_number}/{len(reader.pages)}"
    )


# ============================================================
# VARIABLES
# ============================================================

records = []

current_zone = ""
current_region = ""
current_division = ""


# ============================================================
# REGULAR EXPRESSIONS
# ============================================================

participant_pattern = re.compile(
    r"^\s*(\d+)\s+(.+?)\s+(\+?\d[\d\s\-]{7,})\s+(Paid|Pending)\s*$",
    re.IGNORECASE
)

zone_pattern = re.compile(
    r"^(.*ZONE)$",
    re.IGNORECASE
)


# ============================================================
# PROCESS PDF
# ============================================================

lines = all_text.splitlines()

for line in lines:

    line = line.strip()

    if not line:
        continue

    # Ignore report headings
    if "DEEPER LIFE CAMPUS FELLOWSHIP" in line.upper():
        continue

    if "ELEVATION 2.0" in line.upper():
        continue

    if "PARTICIPANTS BY DIVISION" in line.upper():
        continue

    if "ORGANISED BY" in line.upper():
        continue

    if "ORGANIZED BY" in line.upper():
        continue

    if "S/N NAME TELEPHONE STATUS" in line.upper():
        continue


    # ========================================================
    # ZONE
    # ========================================================

    zone_match = zone_pattern.match(line)

    if zone_match:

        current_zone = line
        current_region = ""
        current_division = ""

        continue


    # ========================================================
    # PARTICIPANT
    # ========================================================

    participant_match = participant_pattern.match(line)

    if participant_match:

        serial_number = participant_match.group(1)
        name = participant_match.group(2).strip()
        telephone = participant_match.group(3).strip()
        status = participant_match.group(4).strip().title()

        records.append({
            "Zone": current_zone,
            "Region": current_region,
            "Division": current_division,
            "S/N": int(serial_number),
            "Name": name,
            "Telephone": telephone,
            "Status": status
        })

        continue


    # ========================================================
    # LOCATION HEADINGS
    # ========================================================

    if not any(char.isdigit() for char in line):

        ignored_words = [
            "STATUS",
            "TELEPHONE",
            "NAME",
            "S/N",
            "PAID",
            "PENDING",
            "REGISTERED"
        ]

        if line.upper() not in ignored_words:

            if not current_region:
                current_region = line

            else:
                current_division = line


# ============================================================
# CREATE DATAFRAME
# ============================================================

print("\nCreating Excel data...")

df = pd.DataFrame(records)


# ============================================================
# REMOVE DUPLICATES
# ============================================================

if not df.empty:

    df = df.drop_duplicates(
        subset=[
            "Zone",
            "Region",
            "Division",
            "S/N",
            "Name",
            "Telephone"
        ]
    )


# ============================================================
# SAVE EXCEL
# ============================================================

print(f"Records extracted: {len(df)}")

with pd.ExcelWriter(
    excel_file,
    engine="openpyxl"
) as writer:

    # Main data
    df.to_excel(
        writer,
        sheet_name="Participants",
        index=False
    )


    # ========================================================
    # SUMMARY
    # ========================================================

    paid_count = len(
        df[
            df["Status"].str.lower() == "paid"
        ]
    )

    pending_count = len(
        df[
            df["Status"].str.lower() == "pending"
        ]
    )

    summary = pd.DataFrame({
        "Description": [
            "Total Participants",
            "Paid",
            "Pending"
        ],
        "Count": [
            len(df),
            paid_count,
            pending_count
        ]
    })

    summary.to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )


    # ========================================================
    # ZONE SUMMARY
    # ========================================================

    if not df.empty:

        zone_summary = (
            df.groupby(
                ["Zone", "Status"]
            )
            .size()
            .unstack(fill_value=0)
            .reset_index()
        )

        zone_summary.to_excel(
            writer,
            sheet_name="Zone Summary",
            index=False
        )


# ============================================================
# FORMAT EXCEL
# ============================================================

print("\nFormatting Excel file...")

workbook = load_workbook(excel_file)

for worksheet in workbook.worksheets:

    # Freeze first row
    worksheet.freeze_panes = "A2"

    # Filter
    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    # Bold headers
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    # Column widths
    for column in worksheet.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            try:

                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:
                    max_length = cell_length

            except:
                pass

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40
        )


workbook.save(excel_file)


# ============================================================
# FINISHED
# ============================================================

print("\n======================================")
print("CONVERSION COMPLETED")
print("======================================")
print(f"Records: {len(df)}")
print(f"Paid: {paid_count}")
print(f"Pending: {pending_count}")
print(f"Excel file: {excel_file}")
print("======================================")


messagebox.showinfo(
    "Completed",
    f"PDF successfully converted to Excel!\n\n"
    f"Records: {len(df)}\n"
    f"Paid: {paid_count}\n"
    f"Pending: {pending_count}"
)